"""Offline tests for the TOR draft+compliance pipeline (stub LLM, real regulation corpus)."""
import json
import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.tor.checks import CHECKS, load_rules, run_compliance
from app.tor.drafting import BOQLine, allocate_budget
from app.tor.graph import build_tor_compliance_graph
from app.tor.procurement import resolve_method
from app.tor.regulations import get_regulation_index
from app.tor.schemas import (
    ChapterKind, DraftChapter, FindingStatus, Installment, ProcurementMethod, ProcurementType,
    TemplateChapter, TORFacts, TORRequest, TORTemplateSpec,
)

# Offline suite: never reach the network. LLM-judge tests use stub LLMs; TestJevJudge
# switches to the jev backend with a fake decisions endpoint.
from app.config import settings as _settings
_settings.TOR_JUDGE_BACKEND = "llm"

TEMPLATE = TORTemplateSpec(template_id="ebidding_service_over_500k", chapters=[
    TemplateChapter(id="background", title="ความเป็นมา", kind=ChapterKind.TEXT),
    TemplateChapter(id="spec", title="คุณลักษณะเฉพาะ", kind=ChapterKind.LIST),
    TemplateChapter(id="budget", title="วงเงินงบประมาณ", kind=ChapterKind.BOQ),
    TemplateChapter(id="payment", title="งวดงานและการจ่ายเงิน", kind=ChapterKind.PAYMENT_SCHEDULE),
    TemplateChapter(id="penalty", title="ค่าปรับและหลักประกัน", kind=ChapterKind.TEXT),
])


def make_request(**kw) -> TORRequest:
    base = dict(project_name="โครงการพัฒนาระบบสืบค้นเอกสาร", agency_name="สำนักงานทดสอบ",
                budget=3_500_000, duration_days=240, raw_requirements="ระบบค้นหาเอกสารราชการ",
                procurement_type=ProcurementType.SERVICES, template=TEMPLATE)
    base.update(kw)
    return TORRequest(**base)


def ch(id, kind, **kw) -> DraftChapter:
    return DraftChapter(id=id, title=id, kind=kind, **kw)


def run_rule(rule_id, req, chapters, facts=TORFacts(), method=ProcurementMethod.E_BIDDING, llm=None):
    registry = load_rules()
    rule = next(r for r in registry.rules if r.id == rule_id)
    sub = registry.model_copy(update={"rules": [rule]})
    return run_compliance(req, method, chapters, facts, get_regulation_index(), llm=llm, registry=sub)


class TestRegistry(unittest.TestCase):
    def test_all_rules_resolve_checks_and_citations(self):
        index = get_regulation_index()
        for r in load_rules().rules:
            self.assertIn(r.check, CHECKS, r.id)
            if r.citation:
                self.assertIsNotNone(index.get_clause(r.citation["doc"], r.citation["clause"]),
                                     f"{r.id}: citation not found in corpus")


class TestDeterministic(unittest.TestCase):
    def test_budget_allocation_is_exact(self):
        items = allocate_budget(1_234_567.89, [BOQLine(name=n, weight_pct=w) for n, w in
                                               [("a", 33.3), ("b", 33.3), ("c", 33.4)]])
        self.assertEqual(round(sum(i.total_price for i in items), 2), 1_234_567.89)

    def test_installments_over_100_fail(self):
        req = make_request()
        bad = ch("payment", ChapterKind.PAYMENT_SCHEDULE, installments=[
            Installment(no=1, percent=60, deliverable="x", due_day=60),
            Installment(no=2, percent=50, deliverable="y", due_day=300)])
        f = run_rule("INSTALLMENTS_SUM_100", req, [bad])
        self.assertEqual(f[0].status, FindingStatus.FAIL)
        f = run_rule("INSTALLMENTS_WITHIN_DURATION", req, [bad])
        self.assertEqual(f[0].status, FindingStatus.FAIL)

    def test_penalty_out_of_range(self):
        f = run_rule("PENALTY_RATE_RANGE", make_request(), [], TORFacts(penalty_rate_pct_per_day=0.5))
        self.assertEqual(f[0].status, FindingStatus.FAIL)
        self.assertIn("ข้อ ๑๖๒", f[0].citation.text)

    def test_penalty_missing_warns_not_passes(self):
        f = run_rule("PENALTY_RATE_RANGE", make_request(), [], TORFacts())
        self.assertEqual(f[0].status, FindingStatus.WARN)

    def test_missing_fact_is_never_a_silent_pass(self):
        f = run_rule("PERFORMANCE_BOND_RATE", make_request(), [], TORFacts())
        self.assertEqual(f[0].status, FindingStatus.WARN)
        f = run_rule("PERFORMANCE_BOND_RATE", make_request(), [], TORFacts(extraction_ok=False))
        self.assertEqual(f[0].status, FindingStatus.NEEDS_HUMAN)

    def test_regex_backstop_when_llm_extractor_fails(self):
        from app.tor.drafting import extract_facts
        terms = ch("terms", ChapterKind.TEXT, text="หลักประกันสัญญาร้อยละ ๓ ค่าปรับเป็นรายวันในอัตราร้อยละ ๐.๕")
        facts = extract_facts([terms], llm=lambda m: "")
        self.assertFalse(facts.extraction_ok)
        self.assertEqual((facts.performance_bond_pct, facts.penalty_rate_pct_per_day), (3.0, 0.5))

    def test_bond_below_standard_fails(self):
        f = run_rule("PERFORMANCE_BOND_RATE", make_request(), [], TORFacts(performance_bond_pct=3))
        self.assertEqual(f[0].status, FindingStatus.FAIL)

    def test_ebidding_wording_in_specific_method(self):
        text = ch("q", ChapterKind.LIST, items=["เป็นนิติบุคคลผู้มีอาชีพรับจ้างงานที่ประกวดราคาอิเล็กทรอนิกส์ดังกล่าว"])
        f = run_rule("METHOD_TERMINOLOGY_CONSISTENT", make_request(budget=300_000), [text],
                     method=ProcurementMethod.SPECIFIC)
        self.assertEqual(f[0].status, FindingStatus.FAIL)

    def test_findings_are_anchored_for_highlighting(self):
        item = "เป็นนิติบุคคลผู้มีอาชีพรับจ้างงานที่ประกวดราคาอิเล็กทรอนิกส์ดังกล่าว"
        text = ch("q", ChapterKind.LIST, items=["มีความสามารถตามกฎหมาย", item])
        f = run_rule("METHOD_TERMINOLOGY_CONSISTENT", make_request(budget=300_000), [text],
                     method=ProcurementMethod.SPECIFIC)[0]
        self.assertEqual(f.anchor.item_index, 1)
        self.assertEqual(item[f.anchor.start:f.anchor.end], "ประกวดราคาอิเล็กทรอนิกส์")

    def test_unverified_rule_never_fails(self):
        f = run_rule("REFERENCE_WORK_MAX_50PCT", make_request(), [],
                     TORFacts(min_reference_work_value=2_800_000))
        self.assertEqual(f[0].status, FindingStatus.NEEDS_HUMAN)

    def test_method_resolution(self):
        self.assertEqual(resolve_method(make_request(budget=300_000))[0], ProcurementMethod.SPECIFIC)
        self.assertEqual(resolve_method(make_request())[0], ProcurementMethod.E_BIDDING)
        self.assertEqual(resolve_method(make_request(procurement_method=ProcurementMethod.SELECTION))[0],
                         ProcurementMethod.SELECTION)


class TestJudgeGrounding(unittest.TestCase):
    spec = ch("spec", ChapterKind.LIST, items=["เครื่องแม่ข่าย Dell PowerEdge R760 จำนวน 2 เครื่อง"])

    def _judge(self, quote):
        reply = json.dumps({"violations": [{"chapter_id": "spec", "quote": quote,
                                            "reason": "ระบุยี่ห้อ", "suggestion": "ระบุคุณลักษณะแทน"}]},
                           ensure_ascii=False)
        return run_rule("SPEC_NO_BRAND_LOCK", make_request(), [self.spec], llm=lambda m: reply)

    def test_grounded_violation_fails(self):
        f = self._judge("Dell PowerEdge R760")
        self.assertEqual(f[0].status, FindingStatus.FAIL)

    def test_hallucinated_quote_is_not_trusted(self):
        f = self._judge("HPE ProLiant DL380")
        self.assertEqual(f[0].status, FindingStatus.NEEDS_HUMAN)

    def test_llm_down_needs_human(self):
        f = run_rule("SPEC_NO_BRAND_LOCK", make_request(), [self.spec], llm=lambda m: "")
        self.assertEqual(f[0].status, FindingStatus.NEEDS_HUMAN)


class ScriptedLLM:
    """Writes a flawed first draft (brand lock, 110% installments, 0.5% penalty), then fixes on revision."""

    def __call__(self, messages):
        p = messages[0]["content"]
        revising = "ร่างเดิมไม่ผ่าน" in p
        if p.startswith("คุณเป็นผู้ตรวจสอบ") and "ตอบทุกข้อของรายการตรวจ" in p:  # checklist-style judge
            return '{"answers": []}'
        if p.startswith("คุณเป็นผู้ตรวจสอบ"):          # judge
            issue = p.split("ตรวจเฉพาะประเด็นนี้เท่านั้น:")[1].split("\n")[0]
            if "Cisco" in p and "ยี่ห้อ" in issue:
                return json.dumps({"violations": [{"chapter_id": "spec", "quote": "อุปกรณ์ Cisco Catalyst 9300",
                                                   "reason": "ระบุยี่ห้อ", "suggestion": "สวิตช์ 48 พอร์ต หรือเทียบเท่า"}]},
                                  ensure_ascii=False)
            return '{"violations": []}'
        if p.startswith("สกัดข้อมูล"):               # fact extraction
            rate = 0.1 if "ร้อยละ 0.1 ต่อวัน" in p else 0.5
            return json.dumps({"penalty_rate_pct_per_day": rate, "performance_bond_pct": 5,
                               "subcontract_penalty_pct": 10})
        chapter = p.split("บทที่ต้องร่าง:")[1].split("\n")[0].strip()
        if chapter == "ความเป็นมา":
            return json.dumps({"text": "หน่วยงานมีความจำเป็นต้องพัฒนาระบบสืบค้นเอกสารราชการเพื่อเพิ่มประสิทธิภาพการให้บริการ"},
                              ensure_ascii=False)
        if chapter == "คุณลักษณะเฉพาะ":
            item = "สวิตช์เครือข่ายขนาด 48 พอร์ต หรือเทียบเท่า" if revising else "อุปกรณ์ Cisco Catalyst 9300"
            return json.dumps({"items": [item, "รองรับการค้นหาภาษาไทย"]}, ensure_ascii=False)
        if chapter == "วงเงินงบประมาณ":
            return json.dumps({"items": [{"name": "พัฒนาระบบ", "weight_pct": 70}, {"name": "อบรม", "weight_pct": 30}]},
                              ensure_ascii=False)
        if chapter == "งวดงานและการจ่ายเงิน":
            second = 60 if revising else 70
            return json.dumps({"installments": [{"no": 1, "percent": 40, "deliverable": "แผนงาน", "due_day": 60},
                                                {"no": 2, "percent": second, "deliverable": "ระบบ", "due_day": 240}]},
                              ensure_ascii=False)
        if chapter == "ค่าปรับและหลักประกัน":
            rate = "0.1" if revising else "0.5"
            return json.dumps({"text": f"กำหนดค่าปรับร้อยละ {rate} ต่อวัน หลักประกันสัญญาร้อยละ 5 ห้ามจ้างช่วงโดยไม่ได้รับอนุญาต ปรับร้อยละ 10"},
                              ensure_ascii=False)
        return ""


class TestEndToEnd(unittest.TestCase):
    def test_flawed_draft_is_revised_until_compliant(self):
        graph = build_tor_compliance_graph(llm=ScriptedLLM(), max_revisions=2)
        result = graph.invoke({"request": make_request()})["result"]
        report = result.report
        self.assertGreaterEqual(report.revisions, 1)
        self.assertTrue(report.passed, [f.model_dump() for f in report.blocking])
        spec = next(c for c in result.chapters if c.id == "spec")
        self.assertNotIn("Cisco", spec.as_plain_text())
        self.assertEqual(sum(i.percent for i in next(c for c in result.chapters if c.id == "payment").installments), 100)
        self.assertEqual(result.facts.penalty_rate_pct_per_day, 0.1)

    def test_revision_budget_is_bounded(self):
        class NeverFixes(ScriptedLLM):
            def __call__(self, messages):
                messages[0]["content"] = messages[0]["content"].replace("ร่างเดิมไม่ผ่าน", "")
                return super().__call__(messages)
        result = build_tor_compliance_graph(llm=NeverFixes(), max_revisions=2).invoke(
            {"request": make_request()})["result"]
        self.assertEqual(result.report.revisions, 2)
        self.assertFalse(result.report.passed)


class TestPostGeneration(unittest.TestCase):
    """revise_chapter / apply_user_edit / acknowledge / export_gate with a stub LLM."""

    def setUp(self):
        import tempfile
        from app.tor.service import AuditLog
        self.audit = AuditLog(tempfile.mkdtemp())
        self.base = build_tor_compliance_graph(llm=ScriptedLLM()).invoke({"request": make_request()})["result"]

    def test_user_brand_instruction_is_obeyed_and_flagged(self):
        from app.tor.service import revise_chapter, export_gate
        class UserWantsCisco(ScriptedLLM):
            def __call__(self, messages):
                p = messages[0]["content"]
                if "คำสั่งแก้ไขจากผู้ใช้" in p:
                    return json.dumps({"items": ["อุปกรณ์ Cisco Catalyst 9300", "รองรับการค้นหาภาษาไทย"]},
                                      ensure_ascii=False)
                return super().__call__(messages)
        out = revise_chapter(self.base, "spec", "ระบุ Cisco Catalyst 9300", actor="u1",
                             llm=UserWantsCisco(), audit=self.audit)
        self.assertIn("Cisco", next(c for c in out.chapters if c.id == "spec").as_plain_text())
        self.assertEqual([c for c in out.chapters if c.id != "spec"],
                         [c for c in self.base.chapters if c.id != "spec"])
        self.assertEqual(out.version, self.base.version + 1)
        self.assertIn("SPEC_NO_BRAND_LOCK:spec", [f.key for f in export_gate(out)])
        self.assertEqual(self.audit.read(out.document_id)[-1]["event"], "ai_revision")

    def test_acknowledge_unblocks_export_and_survives_only_while_finding_exists(self):
        from app.tor.service import apply_user_edit, acknowledge, export_gate
        bad = next(c for c in self.base.chapters if c.id == "payment").model_copy(update={"installments": [
            Installment(no=1, percent=50, deliverable="x", due_day=60),
            Installment(no=2, percent=60, deliverable="y", due_day=200)]})
        out = apply_user_edit(self.base, bad, actor="u1", llm=ScriptedLLM(), audit=self.audit)
        self.assertIn("INSTALLMENTS_SUM_100:payment", [f.key for f in export_gate(out)])
        with self.assertRaises(ValueError):
            acknowledge(out, ["INSTALLMENTS_SUM_100:payment"], "u1", "  ", audit=self.audit)
        out = acknowledge(out, ["INSTALLMENTS_SUM_100:payment"], "u1", "ผู้บริหารอนุมัติแล้ว", audit=self.audit)
        self.assertEqual(export_gate(out), [])
        fixed = apply_user_edit(out, next(c for c in self.base.chapters if c.id == "payment"), "u1",
                                llm=ScriptedLLM(), audit=self.audit)
        self.assertEqual(fixed.report.acknowledgements, [])
        self.assertEqual([e["event"] for e in self.audit.read(out.document_id)],
                         ["user_edit", "acknowledge", "user_edit"])



class TestConcurrency(unittest.TestCase):
    def test_in_flight_llm_calls_never_exceed_limit(self):
        import threading, time
        from app.config import settings
        lock, state = threading.Lock(), {"now": 0, "peak": 0}

        class Slow(ScriptedLLM):
            def __call__(self, messages):
                with lock:
                    state["now"] += 1
                    state["peak"] = max(state["peak"], state["now"])
                time.sleep(0.05)
                try:
                    return super().__call__(messages)
                finally:
                    with lock:
                        state["now"] -= 1
        build_tor_compliance_graph(llm=Slow()).invoke({"request": make_request()})
        self.assertLessEqual(state["peak"], settings.TOR_LLM_CONCURRENCY)
        self.assertGreater(state["peak"], 1, "chapters should actually run in parallel")

    def test_unchanged_chapters_are_not_rejudged(self):
        from app.tor.checks import verify
        from app.tor.concurrency import JudgeCache
        calls = {"judge": 0}

        class Counting(ScriptedLLM):
            def __call__(self, messages):
                if messages[0]["content"].startswith("คุณเป็นผู้ตรวจสอบ"):
                    calls["judge"] += 1
                return super().__call__(messages)
        llm, cache = Counting(), JudgeCache()
        req = make_request()
        chapters = build_tor_compliance_graph(llm=ScriptedLLM()).invoke({"request": req})["result"].chapters
        verify(req, ProcurementMethod.E_BIDDING, chapters, get_regulation_index(), llm, cache)
        first = calls["judge"]
        edited = [c.model_copy(update={"text": c.text + " (แก้ไข)"}) if c.id == "background" else c for c in chapters]
        verify(req, ProcurementMethod.E_BIDDING, edited, get_regulation_index(), llm, cache)
        self.assertGreater(first, 1)
        # Both judge rules cover 'background' (no qualification-titled chapter in this template),
        # so the edited chapter is re-judged once per rule; every other chapter comes from cache.
        self.assertEqual(calls["judge"] - first, 2, "only the edited chapter should be re-judged")



class TestJevJudge(unittest.TestCase):
    """Jev backend with a fake decisions endpoint (keyword-driven probabilities)."""

    def setUp(self):
        from app.config import settings
        from app.services import jev_client
        self.settings, self.jev_client = settings, jev_client
        self._backend, self._decide = settings.TOR_JUDGE_BACKEND, jev_client.decide
        settings.TOR_JUDGE_BACKEND = "jev"
        self.calls = 0

        def fake(state, questions):
            self.calls += 1
            out = {}
            for qid, q in questions.items():
                text = q["instructions"]
                p = 0.95 if "Dell" in text else 0.55 if "Lenovo" in text else 0.02
                other = next(k for k in q["criteria"] if k != "compliant")
                out[qid] = {"type": "choice", "probabilities": {other: p, "compliant": round(1 - p, 4)}}
            return out
        jev_client.decide = fake

    def tearDown(self):
        self.settings.TOR_JUDGE_BACKEND = self._backend
        self.jev_client.decide = self._decide

    def spec(self, *items):
        return ch("spec", ChapterKind.LIST, items=list(items))

    def test_thresholds_map_to_fail_review_pass(self):
        f = run_rule("SPEC_NO_BRAND_LOCK", make_request(),
                     [self.spec("เครื่องแม่ข่าย Dell R760", "โน้ตบุ๊ก Lenovo หรือเทียบเท่า", "รองรับผู้ใช้ 500 ราย")])
        by_status = {x.status: x for x in f}
        self.assertEqual(by_status[FindingStatus.FAIL].evidence, "เครื่องแม่ข่าย Dell R760")
        self.assertEqual(by_status[FindingStatus.NEEDS_HUMAN].evidence, "โน้ตบุ๊ก Lenovo หรือเทียบเท่า")
        self.assertEqual(len(f), 2, "the compliant unit must not produce a finding")
        self.assertAlmostEqual(by_status[FindingStatus.FAIL].confidence, 0.95)

    def test_clean_chapter_passes(self):
        f = run_rule("SPEC_NO_BRAND_LOCK", make_request(), [self.spec("รองรับผู้ใช้ 500 ราย")])
        self.assertEqual([x.status for x in f], [FindingStatus.PASS])

    def test_jev_outage_needs_human(self):
        def down(state, questions):
            raise RuntimeError("503")
        self.jev_client.decide = down
        f = run_rule("SPEC_NO_BRAND_LOCK", make_request(), [self.spec("รองรับผู้ใช้ 500 ราย")])
        self.assertEqual(f[0].status, FindingStatus.NEEDS_HUMAN)

    def test_checklist_rule_uses_items_as_criteria(self):
        seen = {}
        def spy(state, questions):
            seen.update(next(iter(questions.values()))["criteria"])
            return {qid: {"probabilities": {"compliant": 1.0}} for qid in questions}
        self.jev_client.decide = spy
        quals = ch("q", ChapterKind.LIST, items=["มีความสามารถตามกฎหมาย"])
        quals.title = "คุณสมบัติของผู้ยื่นข้อเสนอ"
        run_rule("QUALIFICATIONS_PROPORTIONATE", make_request(), [quals])
        self.assertIn("item_1", seen)
        self.assertIn("compliant", seen)

    def test_prose_is_split_into_sentence_units(self):
        from app.tor.jev_judge import split_units
        prose = ch("t", ChapterKind.TEXT, text="ประโยคแรกที่ยาวพอสมควรสำหรับการตรวจสอบรายละเอียด.  ประโยคที่สองที่ยาวพอสมควรเช่นกันสำหรับการตรวจ")
        self.assertEqual(len(split_units(prose)), 2)

    def test_unchanged_units_hit_cache(self):
        from app.tor.checks import verify
        from app.tor.concurrency import JudgeCache
        cache, req = JudgeCache(), make_request()
        chapters = [self.spec("รองรับผู้ใช้ 500 ราย"), ch("background", ChapterKind.TEXT, text="ความเป็นมาของโครงการโดยละเอียด")]
        verify(req, ProcurementMethod.E_BIDDING, chapters, get_regulation_index(), ScriptedLLM(), cache)
        first = self.calls
        verify(req, ProcurementMethod.E_BIDDING, chapters, get_regulation_index(), ScriptedLLM(), cache)
        self.assertGreater(first, 0)
        self.assertEqual(self.calls, first, "second verify of identical chapters must not call Jev")



class TestWebAndExport(unittest.TestCase):
    def setUp(self):
        import tempfile
        from app.tor.service import AuditLog
        self.tmp = tempfile.mkdtemp()
        self.audit = AuditLog(self.tmp)
        self.result = build_tor_compliance_graph(llm=ScriptedLLM()).invoke({"request": make_request()})["result"]

    def test_thai_text_gets_invisible_word_breaks_only_between_thai_words(self):
        from app.tor.export_docx import breakable
        out = breakable("ระบบต้องรองรับ API จำนวน 500 ราย")
        self.assertEqual(out.replace("​", ""), "ระบบต้องรองรับ API จำนวน 500 ราย")
        self.assertIn("ระบบ​ต้อง", out)
        self.assertNotIn("API​", out)

    def test_view_is_json_and_hides_internals(self):
        from app.tor.api import to_view
        view = to_view(self.result)
        payload = json.loads(view.model_dump_json())
        self.assertNotIn("request", payload)
        self.assertEqual(payload["project"]["procurement_method"], "e_bidding")
        self.assertTrue(all(f["status"] != "pass" for f in payload["report"]["findings"]))
        self.assertTrue(payload["report"]["export_ready"])

    def test_export_blocked_until_acknowledged_and_draft_allowed(self):
        import docx
        from app.tor.export_docx import ExportBlocked
        from app.tor.service import acknowledge, apply_user_edit, export_docx
        bad = next(c for c in self.result.chapters if c.id == "payment").model_copy(update={"installments": [
            Installment(no=1, percent=50, deliverable="x", due_day=60),
            Installment(no=2, percent=60, deliverable="y", due_day=200)]})
        res = apply_user_edit(self.result, bad, actor="u1", llm=ScriptedLLM(), audit=self.audit)
        out = os.path.join(self.tmp, "tor.docx")
        with self.assertRaises(ExportBlocked):
            export_docx(res, out, actor="u1", audit=self.audit)
        draft = export_docx(res, out, actor="u1", draft=True, audit=self.audit)
        header = docx.Document(draft).sections[0].header.paragraphs[0].text
        self.assertIn("ฉบับร่าง", header)
        res = acknowledge(res, ["INSTALLMENTS_SUM_100:payment"], "u1", "อนุมัติแล้ว", audit=self.audit)
        final = export_docx(res, out, actor="u1", audit=self.audit)
        d = docx.Document(final)
        self.assertEqual(d.sections[0].header.paragraphs[0].text, "")
        body = "\n".join(p.text for p in d.paragraphs).replace("​", "")  # strip renderer word breaks
        self.assertIn("ร่างขอบเขตของงาน", body)
        self.assertIn("1. ความเป็นมา", body)
        self.assertEqual(len(d.tables), 2, "BOQ and payment tables")
        self.assertIn("สามล้านห้าแสนบาทถ้วน", d.tables[0].rows[-1].cells[0].text.replace("​", ""))
        self.assertEqual([e["event"] for e in self.audit.read(res.document_id)][-2:], ["acknowledge", "export"])


if __name__ == "__main__":
    unittest.main()
