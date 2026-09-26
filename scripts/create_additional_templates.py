import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def set_thai_font(run, font_name="TH Sarabun PSK", size_pt=16, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold

def create_meeting_agenda_template(output_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.98)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(1.18)
        s.right_margin = Inches(0.79)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("ระเบียบวาระการประชุม{{ committee_name }}\n")
    set_thai_font(r, size_pt=18, bold=True)
    r = p_title.add_run("ครั้งที่ {{ meeting_no }}/{{ meeting_year }}\n")
    set_thai_font(r, size_pt=16, bold=True)
    r = p_title.add_run("วัน{{ meeting_date }} เวลา {{ meeting_time }} น.\nณ {{ meeting_location }}\n")
    set_thai_font(r, size_pt=16)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5 standard agendas
    agendas = [
        ("ระเบียบวาระที่ 1", "เรื่องที่ประธานแจ้งให้ที่ประชุมทราบ", "{{ agenda_1_chairman_notes }}"),
        ("ระเบียบวาระที่ 2", "เรื่องรับรองรายงานการประชุมครั้งที่ผ่านมา", "{{ agenda_2_previous_minutes }}"),
        ("ระเบียบวาระที่ 3", "เรื่องสืบเนื่อง / เพื่อทราบ", "{% for item in agenda_3_matters_to_inform %}\n  3.{{ loop.index }} {{ item }}\n{% endfor %}"),
        ("ระเบียบวาระที่ 4", "เรื่องเพื่อพิจารณา", "{% for item in agenda_4_matters_for_consideration %}\n  4.{{ loop.index }} {{ item }}\n{% endfor %}"),
        ("ระเบียบวาระที่ 5", "เรื่องอื่นๆ (ถ้ามี)", "{{ agenda_5_other_matters | default('ไม่มี') }}")
    ]

    for no, title, content in agendas:
        p_h = doc.add_paragraph()
        r = p_h.add_run(f"{no}: {title}")
        set_thai_font(r, size_pt=16, bold=True)
        p_c = doc.add_paragraph()
        p_c.paragraph_format.first_line_indent = Inches(0.5)
        r = p_c.add_run(content)
        set_thai_font(r, size_pt=16)

    doc.save(output_path)
    print(f"Created Meeting Agenda template: {output_path}")

def create_leave_template(output_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.98)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(1.18)
        s.right_margin = Inches(0.79)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("แบบใบลาป่วย ลากิจส่วนตัว\n")
    set_thai_font(r, size_pt=18, bold=True)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_meta.add_run("เขียนที่ {{ write_location }}\nวันที่ {{ doc_date }}\n")
    set_thai_font(r, size_pt=16)

    p_sub = doc.add_paragraph()
    r = p_sub.add_run("เรื่อง: ")
    set_thai_font(r, size_pt=16, bold=True)
    r = p_sub.add_run("ขอลา{{ leave_type }}\n")
    set_thai_font(r, size_pt=16)
    r = p_sub.add_run("เรียน: ")
    set_thai_font(r, size_pt=16, bold=True)
    r = p_sub.add_run("{{ supervisor_title }}\n")
    set_thai_font(r, size_pt=16)

    p_body = doc.add_paragraph()
    p_body.paragraph_format.first_line_indent = Inches(0.5)
    r = p_body.add_run(
        "ข้าพเจ้า {{ employee_name }} ตำแหน่ง {{ employee_position }} สังกัด {{ department_name }}\n"
        "มีความประสงค์ขอลา{{ leave_type }} เนื่องจาก {{ leave_reason }} "
        "ตั้งแต่วันที่ {{ start_date }} ถึงวันที่ {{ end_date }} กำหนด {{ total_days }} วัน\n"
        "ในระหว่างลาจะสามารถติดต่อข้าพเจ้าได้ที่ {{ contact_address_and_phone }}"
    )
    set_thai_font(r, size_pt=16)

    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(24)
    r = p_sign.add_run(
        "(ลงชื่อ)............................................................\n"
        "({{ employee_name }})\n"
        "ตำแหน่ง {{ employee_position }}"
    )
    set_thai_font(r, size_pt=16)

    doc.save(output_path)
    print(f"Created Leave template: {output_path}")

if __name__ == '__main__':
    create_meeting_agenda_template('app/templates/meeting_agenda_template.docx')
    create_leave_template('app/templates/leave_template.docx')
